"""Tables: Excel / Word objects embedded in the drawing, AutoCAD tables, tables drawn as aligned
texts, and the setting-out coordinates they hold (footprints of objects that may not be drawn)."""

from __future__ import annotations

import io
import re
import zipfile
from itertools import pairwise

from ezdxf.document import Drawing as DxfDocument
from shapely.geometry import Polygon

from .entities import text_lines
from .model import EmbeddedTable


def _zip_from_ole(data: bytes) -> zipfile.ZipFile | None:
    i = data.find(b"PK\x03\x04")
    j = data.rfind(b"PK\x05\x06")
    if i < 0 or j < 0:
        return None
    try:
        return zipfile.ZipFile(io.BytesIO(data[i : j + 22]))
    except Exception:
        return None


def _parse_setting_out(rows: list[list]) -> dict[str, list[tuple[str, float, float]]]:
    """Recognise setting-out tables: an 'Ob…' header row naming the object, then rows of
    point name + two coordinates."""
    objects: dict[str, list] = {}
    current = None
    for r in rows:
        vals = [v for v in r if v not in (None, "")]
        if not vals:
            continue
        strs = [v for v in vals if isinstance(v, str)]
        nums = [v for v in vals if isinstance(v, (int, float))]
        if len(vals) >= 2 and isinstance(vals[0], str) and vals[0].strip().lower().startswith("ob") and not nums:
            current = str(vals[-1]).strip()
            objects.setdefault(current, [])
        elif current and len(nums) >= 2 and strs:
            objects[current].append((strs[0].strip(), float(nums[0]), float(nums[1])))
    return {k: v for k, v in objects.items() if v}


def extract_tables(doc: DxfDocument) -> list[EmbeddedTable]:
    tables: dict[str, EmbeddedTable] = {}
    for name in doc.layouts.names():
        layout = doc.layouts.get(name)
        for e in layout.query("OLE2FRAME"):
            try:
                z = _zip_from_ole(e.binary_data())
            except Exception:
                z = None
            if z is None:
                continue
            names = z.namelist()
            rows: list[list] = []
            kind = ""
            if any(n.startswith("xl/") for n in names):
                kind = "xlsx"
                try:
                    import openpyxl

                    wb = openpyxl.load_workbook(io.BytesIO(_rezip(z)), data_only=True)
                    for ws in wb.worksheets:
                        for row in ws.iter_rows(values_only=True):
                            if any(v is not None for v in row):
                                rows.append(list(row))
                except Exception:
                    continue
            elif "word/document.xml" in names:
                kind = "docx"
                xml = z.read("word/document.xml").decode("utf8", "ignore")
                for tr in re.findall(r"<w:tr[ >].*?</w:tr>", xml, re.S):
                    cells = []
                    for tc in re.findall(r"<w:tc[ >].*?</w:tc>", tr, re.S):
                        cells.append(" ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", tc)).strip())
                    rows.append(cells)
                if not rows:
                    rows = [[" ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))]]
            else:
                continue
            _add_table(tables, name, kind, rows)
        for e in layout.query("ACAD_TABLE"):
            rows = _acad_table_rows(doc, e)
            if rows:
                _add_table(tables, name, "acad_table", rows)
    return list(tables.values())


def _add_table(tables: dict, layout_name: str, kind: str, rows: list[list]) -> None:
    rows = [[v for v in r if v not in (None, "")] for r in rows]  # trim empty cells
    rows = [r for r in rows if r]
    if not rows:
        return
    key = repr(rows)
    if key in tables:
        if layout_name not in tables[key].layouts:
            tables[key].layouts.append(layout_name)
        return
    objects = _parse_setting_out(rows)
    title = ", ".join(objects) if objects else " ".join(str(v) for v in rows[0])[:80]
    tables[key] = EmbeddedTable(f"T{len(tables) + 1}", [layout_name], kind, title, rows, objects)


def _cell_value(s: str):
    t = s.strip().replace("\u202f", "").replace(" ", "")
    if re.fullmatch(r"[-+]?\d+([.,]\d+)?", t):
        try:
            return float(t.replace(",", "."))
        except ValueError:
            pass
    return s.strip()


def _acad_table_rows(doc: DxfDocument, e) -> list[list]:
    """Cell texts of an ACAD_TABLE, read from the anonymous block AutoCAD keeps with its graphics."""
    handle = e.dxf.get("block_record_handle")
    if not handle:
        return []
    rec = doc.entitydb.get(handle)
    blk = doc.blocks.get(rec.dxf.name) if rec is not None else None
    if blk is None:
        return []
    cells = []
    for t in blk:
        if t.dxftype() in ("TEXT", "MTEXT"):
            lines = text_lines(t)
            if lines:
                text = " ".join(x[0] for x in lines)
                y = sum(x[2] for x in lines) / len(lines)
                cells.append((text, lines[0][1], y, lines[0][3]))
    return [[_cell_value(c) for c in r] for r in _rows_of_texts(cells)]


def _rows_of_texts(cells: list[tuple[str, float, float, float]]) -> list[list[str]]:
    """Group (text, x, y, height) into rows (same y within half a text height), left to right."""
    rows: list[list] = []
    for c in sorted(cells, key=lambda c: -c[2]):
        if rows and abs(rows[-1][0][2] - c[2]) <= 0.5 * max(c[3], rows[-1][0][3], 1e-9):
            rows[-1].append(c)
        else:
            rows.append([c])
    return [[c[0] for c in sorted(r, key=lambda c: c[1])] for r in rows]


_NUMBER = re.compile(r"\d")


def text_tables(texts, space: str, max_tables: int = 8) -> list[EmbeddedTable]:
    """Tables drawn as plain texts (e.g. a land-use balance on a sheet): runs of at least three
    rows of two or more aligned texts, with numbers in most rows."""
    lines = [(s, x, y, h) for s, x, y, h, _ in texts if h]
    if len(lines) < 6:
        return []
    rows: list[list] = []
    for c in sorted(lines, key=lambda c: -c[2]):
        if rows and abs(rows[-1][0][2] - c[2]) <= 0.5 * max(c[3], rows[-1][0][3]):
            rows[-1].append(c)
        else:
            rows.append([c])
    out, run = [], []

    def aligned(run) -> bool:
        """Most cells start at a column position shared with other rows (left-aligned columns)."""
        h = max(c[3] for r in run for c in r)
        xs = [c[1] for r in run for c in r]
        shared = sum(
            1
            for r in run
            for c in r
            if sum(1 for r2 in run if r2 is not r and any(abs(c2[1] - c[1]) <= 1.5 * h for c2 in r2)) >= 1
        )
        return shared >= 0.7 * len(xs)

    def flush():
        if (
            len(run) >= 3
            and sum(1 for r in run if any(_NUMBER.search(c[0]) for c in r)) >= len(run) / 2
            and sum(1 for r in run if len(r) >= 2) >= 0.75 * len(run)
            and aligned(run)
        ):
            out.append(
                EmbeddedTable(
                    f"X{len(out) + 1}",
                    [space],
                    "text",
                    " ".join(c[0] for c in sorted(run[0], key=lambda c: c[1]))[:80],
                    [[_cell_value(c[0]) for c in sorted(r, key=lambda c: c[1])] for r in run[:60]],
                )
            )

    for r in rows:
        h = max(c[3] for c in r)
        xs = sorted(c[1] for c in r)
        tabular = len(r) >= 2 and max(b - a for a, b in pairwise(xs)) > 2 * h
        if tabular and run:
            prev = run[-1]
            gap = prev[0][2] - r[0][2]
            px0, px1 = min(c[1] for c in prev), max(c[1] for c in prev)
            if gap <= 3.5 * h and xs[0] <= px1 and xs[-1] >= px0:
                run.append(r)
                continue
        flush()
        run = [r] if tabular else []
        if len(out) >= max_tables:
            return out
    flush()
    return out[:max_tables]


def _rezip(z: zipfile.ZipFile) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as out:
        for n in z.namelist():
            out.writestr(n, z.read(n))
    return buf.getvalue()


def table_polygon(points: list[tuple[str, float, float]]) -> Polygon:
    """Footprint polygon from setting-out points (centre points such as 'M1.C' are ignored)."""
    pts = [(e, n) for name, n, e in points if not name.strip().upper().endswith(".C")]
    if len(pts) < 3:
        return Polygon()
    poly = Polygon(pts)
    if not poly.is_valid or poly.area <= 0:
        poly = poly.convex_hull
    return poly
