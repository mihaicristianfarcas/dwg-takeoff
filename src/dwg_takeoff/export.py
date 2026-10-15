"""Excel export of a takeoff result (summary, per-entity evidence, reasoning, checks)."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS = {"measured": "măsurat din planșă", "assumption_based": "cu ipoteze", "not_in_drawing": "lipsă în planșă"}
CONF = {"high": "mare", "medium": "medie", "low": "scăzută"}


ORIGIN = {"assumed": "presupus", "document": "din documente", "drawing": "din planșă"}
MAX_FACTORS = 4
FACTOR_VALUE_COLS = [get_column_letter(8 + 2 * k) for k in range(MAX_FACTORS)]  # H, J, L, N
MEASURE_RO = {
    "area": "arie",
    "length": "lungime",
    "perimeter": "perimetru",
    "count": "număr",
    "strip_area": "arie benzi",
    "enclosed_area": "arie închisă",
    "count_groups": "număr grupuri",
}
INPUT_FILL = PatternFill("solid", fgColor="FFF4D6")  # factor values the estimator may change


def r2(x):
    return round(x, 2) if isinstance(x, (int, float)) else x


HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
WRAP = Alignment(wrap_text=True, vertical="top")


def _sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, WRAP
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = "A2"


def result_to_xlsx(res: dict, project_name: str = "", project: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rezumat"
    rows = []
    for i, it in enumerate(res.get("items", []), 1):
        factors = []
        for c in it.get("components", []):
            for f in c.get("factors", []):
                factors.append(
                    f"{f.get('name')}: {f.get('value')} {f.get('unit')} ({ORIGIN.get(f.get('origin'), 'din planșă')})"
                )
        rows.append(
            [
                i,
                it.get("label"),
                r2(it.get("value")),
                it.get("unit"),
                STATUS.get(it.get("status"), it.get("status")),
                CONF.get(it.get("confidence"), it.get("confidence")),
                it.get("formula"),
                "\n".join(factors),
                "\n".join(it.get("assumptions", [])),
                "\n".join(it.get("questions", [])),
            ]
        )
    _sheet(
        ws,
        ["Nr.", "Element", "Cantitate", "UM", "Status", "Încredere", "Formulă", "Factori", "Ipoteze", "Întrebări"],
        rows,
        [5, 30, 13, 6, 16, 10, 50, 40, 60, 50],
    )
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 3).number_format = "#,##0" if ws.cell(r, 4).value == "buc" else "#,##0.00"

    # live calculation: base value × factors as formulas, so a thickness can be changed in Excel
    ws2 = wb.create_sheet("Componente")
    rows = []
    plan_items = {pi.get("id"): pi for pi in (res.get("plan") or {}).get("items", [])}
    for i, it in enumerate(res.get("items", []), 1):
        for ci, c in enumerate(it.get("components", [])):
            pcs = (plan_items.get(it.get("id")) or {}).get("components") or []
            sources = _sources_text(pcs[ci]) if ci < len(pcs) else ""
            st = c.get("stats") or {}
            fs = list(c.get("factors") or [])
            if len(fs) > MAX_FACTORS:  # rare: fold the extra factors into the last column
                extra = 1.0
                for f in fs[MAX_FACTORS - 1 :]:
                    extra *= float(f.get("value") or 0)
                fs = [
                    *fs[: MAX_FACTORS - 1],
                    {
                        "name": " × ".join(f.get("name", "") for f in fs[MAX_FACTORS - 1 :]),
                        "value": extra,
                        "unit": "",
                        "origin": "assumed",
                    },
                ]
            fcells = []
            for k in range(MAX_FACTORS):
                if k < len(fs):
                    f = fs[k]
                    tag = ORIGIN.get(f.get("origin"), "din planșă")
                    fcells += [f"{f.get('name')} ({f.get('unit', '')}, {tag})", f.get("value")]
                else:
                    fcells += [None, None]
            r = len(rows) + 2
            fx = "*".join(f'IF({col}{r}="",1,{col}{r})' for col in FACTOR_VALUE_COLS)
            base = c.get("base_value")
            rows.append(
                [
                    i,
                    it.get("label"),
                    c.get("label"),
                    MEASURE_RO.get(c.get("measure"), c.get("measure")),
                    round(base, 4) if isinstance(base, (int, float)) else base,
                    c.get("base_unit"),
                    *fcells,
                    f"=E{r}*{fx}",
                    c.get("unit"),
                    r2(st.get("raw_sum")),
                    c.get("n_entities"),
                    c.get("formula"),
                    "\n".join(c.get("warnings", [])),
                    sources,
                ]
            )
    headers = ["Nr. element", "Element", "Componentă", "Măsură", "Valoare măsurată", "UM"]
    for k in range(1, MAX_FACTORS + 1):
        headers += [f"Factor {k}", f"Valoare factor {k}"]
    headers += [
        "Cantitate (formulă)",
        "UM rezultat",
        "Suma brută desenată",
        "Nr. entități",
        "Formula aplicației",
        "Avertismente",
        "Surse (catalog / filtre)",
    ]
    _sheet(ws2, headers, rows, [7, 26, 30, 12, 13, 6] + [22, 10] * MAX_FACTORS + [14, 8, 12, 9, 55, 50, 40])
    qcol = get_column_letter(7 + 2 * MAX_FACTORS)
    for r in range(2, ws2.max_row + 1):
        ws2.cell(r, 5).number_format = "#,##0.00"
        ws2.cell(r, 7 + 2 * MAX_FACTORS).number_format = "#,##0.00"
        for k in range(MAX_FACTORS):
            ws2.cell(r, 8 + 2 * k).fill = INPUT_FILL
    # summary quantities = sum of their components (recomputed by Excel when a factor changes)
    for r in range(2, len(res.get("items", [])) + 2):
        it = res["items"][r - 2]
        if it.get("value") is not None and it.get("components"):
            ws.cell(r, 3).value = f"=SUMIF(Componente!$A:$A,A{r},Componente!${qcol}:${qcol})"

    ws3 = wb.create_sheet("Entități")
    rows = []
    for it in res.get("items", []):
        for c in it.get("components", []):
            for e in c.get("evidence", []):
                rows.append(
                    [
                        it.get("label"),
                        c.get("label"),
                        e.get("handle"),
                        e.get("layer"),
                        e.get("kind"),
                        e.get("pattern"),
                        e.get("block"),
                        r2(e.get("raw")),
                        r2(e.get("counted")),
                        e.get("note"),
                        "" if e.get("visible", True) else "strat ascuns",
                    ]
                )
    _sheet(
        ws3,
        [
            "Element",
            "Componentă",
            "Handle AutoCAD",
            "Strat",
            "Tip",
            "Hașură",
            "Bloc / obiect",
            "Valoare entitate",
            "Valoare numărată",
            "Observație",
            "Vizibilitate",
        ],
        rows,
        [26, 28, 12, 34, 10, 14, 22, 14, 14, 44, 12],
    )

    ws4 = wb.create_sheet("Raționament")
    rows = [[it.get("label"), it.get("reasoning")] for it in res.get("items", [])]
    plan = res.get("plan") or {}
    if plan.get("summary"):
        rows.insert(0, ["Rezumat AI", plan["summary"]])
    for n in plan.get("general_notes", []):
        rows.append(["Notă", n])
    for q in plan.get("questions_for_user", []):
        rows.append(["Întrebare", q])
    for ch in res.get("checks", []):
        rows.append(["Verificare", ch.get("message")])
    _sheet(ws4, ["Element", "Explicație"], rows, [28, 120])

    if project_name:
        ws.cell(ws.max_row + 2, 1, f"Planșă: {project_name}")
        if res.get("query"):
            ws.cell(ws.max_row + 1, 1, f"Cerere: {res['query']}")

    # what every number depends on: the drawing, its units, the AI run
    ws5 = wb.create_sheet("Info")
    meta = (res.get("plan") or {}).get("_meta") or {}
    ui = (project or {}).get("unit_info") or {}
    rows = [
        ["Planșă", project_name or (project or {}).get("name", "")],
        ["Cerere", res.get("query") or ""],
        ["Unități desen", (project or {}).get("units", "")],
        [
            "Sursa unităților",
            {
                "header": "antetul desenului",
                "inferred": "deduse din desen",
                "default": "presupuse",
                "override": "alese manual",
            }.get(ui.get("source"), ui.get("source", "")),
        ],
        ["Dovezi unități", "\n".join(ui.get("evidence", []))],
        ["Spațiu măsurat", (project or {}).get("space", "Model")],
        [
            "Motor AI",
            f"{meta.get('backend', '')} · {meta.get('model', '')} · efort {meta.get('effort', '')}".strip(" ·"),
        ],
        ["Durată AI (s)", meta.get("seconds", "")],
        [
            "Informații de proiect",
            (
                f"da, {meta['project_info_chars']} caractere (note / documente de referință; "
                f"textul exact: project_info.txt în rulare)"
            )
            if meta.get("project_info_chars")
            else "nu",
        ],
        ["Verificare AI", _review_text(meta.get("review"))],
        ["Rulare", res.get("run") or ""],
        ["Plan editat manual", "da" if meta.get("edited") else "nu"],
        [
            "Notă",
            "Cantitățile sunt calculate din geometria desenului; AI-ul alege doar sursele și factorii. "
            "Valorile marcate „presupus” trebuie confirmate.",
        ],
        [
            "Recalculare",
            "În foaia Componente, celulele galbene sunt factorii (grosimi, înălțimi...). Modificați-i "
            "și cantitățile din Componente și Rezumat se recalculează.",
        ],
    ]
    _sheet(ws5, ["Câmp", "Valoare"], rows, [22, 110])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _review_text(r: dict | None) -> str:
    if not r:
        return "nu"
    if r.get("error"):
        return "nu a putut fi făcută"
    if r.get("rejected"):
        return "corecturi propuse respinse (nu se potriveau cu desenul)"
    changed = r.get("changed") or []
    head = f"elemente corectate: {', '.join(changed)}" if changed else "planul a fost confirmat"
    return "\n".join([head] + [f"- {n}" for n in r.get("notes") or []])


def _sources_text(comp: dict) -> str:
    """What a component measures: catalogue ids, else the raw filters."""
    out = []
    for key, sign in (("entries", ""), ("subtract_entries", "− ")):
        out += [f"{sign}{e}" for e in comp.get(key) or []]
    for key, sign in (("sources", ""), ("subtract_sources", "− ")):
        for src in comp.get(key) or []:
            if src.get("entry"):
                continue
            bits = [
                src.get("layer") or src.get("table_object") or "",
                src.get("pattern") or "",
                src.get("color") or "",
                src.get("block") or "",
            ]
            out.append(sign + " ".join(b for b in bits if b))
    return ", ".join(out)
