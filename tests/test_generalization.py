"""Drawings that do not look like the client's site plan: no legend on a sheet, everything on
layer 0 told apart by colour, nested blocks, missing or wrong units, geometry drawn on a sheet."""

from types import SimpleNamespace

import ezdxf
import pytest

from dwg_takeoff.analysis.colors import color_name
from dwg_takeoff.analysis.digest import build_digest, digest_to_text
from dwg_takeoff.cad import extract, load_doc
from dwg_takeoff.cad.tables import text_tables
from dwg_takeoff.cad.units import dimension_scale, infer_units
from dwg_takeoff.engine.executor import execute_plan


def rect(x0, y0, x1, y1):
    return [(x0, y0), (x1, y0), (x1, y1), (x0, y1)]


def src(layer, kinds, pattern=None, color=None, block=None):
    return {
        "layer": layer,
        "kinds": kinds,
        "pattern": pattern,
        "color": color,
        "block": block,
        "table_object": None,
        "handles": None,
    }


def item(iid, measure, sources, unit="m²"):
    return {
        "id": iid,
        "label": iid,
        "request_text": "",
        "unit": unit,
        "status": "measured",
        "confidence": "high",
        "reasoning": "",
        "assumptions": [],
        "questions": [],
        "components": [
            {
                "label": iid,
                "measure": measure,
                "sources": sources,
                "subtract_sources": [],
                "region": None,
                "tolerance_m": None,
                "factors": [],
            }
        ],
    }


@pytest.fixture(scope="module")
def garden(tmp_path_factory):
    """Metres, no sheets: decks and lawn told apart only by colour, a tree block placed in rows."""
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 6
    msp = doc.modelspace()
    doc.layers.add("PLAN", color=7)
    # wooden deck: two orange SOLID hatches (ACI 30); lawn: one green hatch by true colour
    for x in (0, 10):
        h = msp.add_hatch(color=30, dxfattribs={"layer": "0"})
        h.paths.add_polyline_path(rect(x, 0, x + 4, 5), is_closed=True)
    lawn = msp.add_hatch(dxfattribs={"layer": "0"})
    lawn.rgb = (60, 160, 60)
    lawn.paths.add_polyline_path(rect(0, 10, 20, 20), is_closed=True)
    # a block whose fill is BYBLOCK: takes the colour of each INSERT
    tile = doc.blocks.new("TILE")
    th = tile.add_hatch(color=0)
    th.paths.add_polyline_path(rect(0, 0, 1, 1), is_closed=True)
    msp.add_blockref("TILE", (30, 0), dxfattribs={"color": 30, "layer": "0"})
    msp.add_blockref("TILE", (32, 0), dxfattribs={"color": 5, "layer": "0"})
    # rows of trees: ROW holds 3 TREE blocks; two rows placed
    tree = doc.blocks.new("TREE")
    tree.add_circle((0, 0), 1)
    row = doc.blocks.new("ROW")
    for i in range(3):
        row.add_blockref("TREE", (i * 5, 0))
    msp.add_blockref("ROW", (0, 30), dxfattribs={"layer": "PLAN"})
    msp.add_blockref("ROW", (0, 40), dxfattribs={"layer": "PLAN"})
    # an MPOLYGON surface (Map 3D / Civil polygons)
    mp = msp.add_mpolygon(dxfattribs={"layer": "PLAN"})
    mp.paths.add_polyline_path(rect(50, 0, 56, 5), is_closed=True)
    # model-space legend: a column of three swatches with labels
    for i, (label, aci) in enumerate((("Deck lemn", 30), ("Gazon", 3), ("Pavaj piatra", 8))):
        sw = msp.add_hatch(color=aci, dxfattribs={"layer": "LEGENDA"})
        sw.paths.add_polyline_path(rect(100, -i * 3, 104, -i * 3 + 2), is_closed=True)
        msp.add_text(label, height=1, dxfattribs={"layer": "LEGENDA"}).set_placement((105, -i * 3 + 0.5))
    path = tmp_path_factory.mktemp("g") / "garden.dxf"
    doc.saveas(path)
    return extract(path, "garden.dxf", load_doc(path))


def test_colours_resolved(garden):
    hatches = [f for f in garden.features if f.kind == "hatch" and not f.nested and f.layer == "0"]
    assert {f.color for f in hatches} == {"#ff7f00", "#3ca03c"}  # ACI 30 orange, true colour green
    tiles = sorted(f.color for f in garden.features if f.nested and f.kind == "hatch")
    assert tiles == ["#0000ff", "#ff7f00"]  # BYBLOCK takes the INSERT colour
    assert color_name("#ff7f00") == "orange" and color_name("#3ca03c") == "green"


def test_area_by_colour(garden):
    plan = {
        "items": [
            item("deck", "area", [src("0", ["hatch"], color="#ff7f00")]),
            item("lawn", "area", [src("0", ["hatch"], color="#3ca03c")]),
        ]
    }
    res = {i["id"]: i for i in execute_plan(plan, garden, {})["items"]}
    assert res["deck"]["value"] == pytest.approx(41)  # 2 decks of 20 m² + the orange tile (1 m²)
    assert res["lawn"]["value"] == pytest.approx(200)


def test_digest_lists_colour_groups(garden):
    text = digest_to_text(build_digest(garden))
    assert "SOLID orange ×2" in text and "SOLID green ×1" in text  # told apart by colour
    assert "block in blocks 'TREE' ×6" in text  # trees placed in "row" blocks


def test_nested_blocks_not_counted_unless_named(garden):
    plan = {
        "items": [
            item("rows", "count", [src("PLAN", ["block"])], "buc"),
            item("trees", "count", [src("PLAN", ["block"], block="TREE")], "buc"),
        ]
    }
    res = {i["id"]: i for i in execute_plan(plan, garden, {})["items"]}
    assert res["rows"]["value"] == 2
    assert res["trees"]["value"] == 6


def test_mpolygon_measured(garden):
    plan = {"items": [item("mp", "area", [src("PLAN", ["hatch"])])]}
    assert execute_plan(plan, garden, {})["items"][0]["value"] == pytest.approx(30)


def test_model_space_legend(garden):
    labels = {(le.layout, le.label, le.color) for le in garden.legends}
    assert ("Model", "Deck lemn", "#ff7f00") in labels
    assert len([le for le in garden.legends if le.layout == "Model"]) == 3


def F(kind, x0, y0, x1, y1, text=None, h=None):
    from shapely.geometry import Point, box

    g = Point(x0, y0) if kind == "text" else box(x0, y0, x1, y1)
    return SimpleNamespace(kind=kind, geom=g, text=text, text_height=h, nested=False, handle=None, block=None)


def test_units_header_kept_when_plausible():
    feats = [F("closed", 0, 0, 40000, 20000), F("closed", 5000, 5000, 6000, 6000), F("text", 100, 100, 0, 0, "A", 250)]
    ui = infer_units(4, feats)
    assert ui.code == 4 and ui.source == "header"


def test_units_feet_saved_as_inches():
    feats = [F("closed", 0, 0, 150, 120)] + [
        F("text", 10 * i, 5, 0, 0, t, 1.5)
        for i, t in enumerate(["50' WIDE PUBLIC R.O.W.", "N16°44'49\"W", "S73°15'11\"E", "25'"])
    ]
    ui = infer_units(1, feats, measurement=0)
    assert ui.code == 2 and ui.source == "inferred"


def test_units_unitless_millimetres():
    feats = [F("closed", 0, 0, 60000, 30000), F("closed", 1000, 1000, 9000, 4000)]
    assert infer_units(0, feats).code == 4
    feats_m = [F("closed", 0, 0, 60, 30), F("closed", 1, 1, 9, 4)]
    assert infer_units(0, feats_m).code == 6


def test_units_from_dimensions():
    # labels show metres with decimals while the lines are 1000x longer: the drawing is in mm
    dims = [(5000.0, 5.0), (12500.0, 12.5), (3200.0, 3.2), (800.0, 0.8)]
    assert dimension_scale(dims)[0] == pytest.approx(0.001)
    feats = [F("closed", 0, 0, 40000, 20000)]
    ui = infer_units(6, feats, dims=dims)
    assert ui.code == 4 and "cote" in " ".join(ui.evidence)
    assert infer_units(0, feats, override=6).code == 6  # the user always has the last word


def test_text_table_detected():
    rows = [("Bilant teritorial", "Suprafata (mp)"), ("Cladire", "400"), ("Alei", "1 610,46"), ("Spatii verzi", "185")]
    texts = []
    for i, (a, b) in enumerate(rows):
        texts.append((a, 0, -i * 5, 2.5, "T"))
        texts.append((b, 60, -i * 5, 2.5, "T"))
    tables = text_tables(texts, "Sheet")
    assert len(tables) == 1 and tables[0].rows[1] == ["Cladire", 400.0]


def test_sheet_only_drawing(tmp_path):
    doc = ezdxf.new("R2018")
    lo = doc.layouts.get("Layout1")
    lo.add_lwpolyline(rect(0, 0, 100, 50), close=True, dxfattribs={"layer": "CONTUR"})
    path = tmp_path / "sheet.dxf"
    doc.saveas(path)
    d = extract(path, "sheet.dxf", load_doc(path))
    assert d.space == "Layout1"
    assert any(f.kind == "closed" for f in d.features)
    assert any("sheet" in w for w in d.warnings)


def test_view_bounds_of_empty_and_point_drawings():
    from dwg_takeoff.web.views import view_bounds as _view_bounds

    empty = SimpleNamespace(focus=None, drawing=SimpleNamespace(features=[], unit_scale=1.0))
    b = _view_bounds(empty)
    assert b[2] > b[0] and b[3] > b[1]
    one = SimpleNamespace(
        focus=None, drawing=SimpleNamespace(features=[F("text", 5, 5, 0, 0, "x", 1)], unit_scale=0.001)
    )
    b = _view_bounds(one)
    assert b[2] - b[0] > 0 and b[3] - b[1] > 0


def test_enclosed_area_from_separate_lines(tmp_path):
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 6
    msp = doc.modelspace()
    pts = [(0, 0), (30, 0), (30, 20), (0, 20)]
    for a, b in zip(pts, pts[1:] + pts[:1], strict=True):
        msp.add_line(a, (b[0] + (0.02 if b == (0, 0) else 0), b[1]), dxfattribs={"layer": "LOT"})  # 2 cm gap
    msp.add_line((40, 0), (50, 0), dxfattribs={"layer": "LOT"})  # stray line, encloses nothing
    path = tmp_path / "lot.dxf"
    doc.saveas(path)
    d = extract(path, "lot.dxf", load_doc(path))
    plan = {"items": [item("lot", "enclosed_area", [src("LOT", ["line"])])]}
    plan["items"][0]["components"][0]["tolerance_m"] = 0.05
    res = execute_plan(plan, d, {})["items"][0]
    assert res["value"] == pytest.approx(600, rel=1e-3)
    notes = [e["note"] for e in res["components"][0]["evidence"]]
    assert sum("liberă" in n for n in notes) == 1


def test_enclosed_area_without_lines(garden):
    plan = {"items": [item("x", "enclosed_area", [src("PLAN", ["block"])])]}
    res = execute_plan(plan, garden, {})["items"][0]
    assert res["value"] == 0 and any("nu închid" in w for w in res["components"][0]["warnings"])


def test_count_groups_of_loose_symbols(tmp_path):
    """Three stairs drawn as treads (lines), one stair inside a block: 3 groups."""
    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 6
    msp = doc.modelspace()
    for x0 in (0, 10, 20):
        for i in range(8):
            msp.add_line((x0, i * 0.3), (x0 + 1.2, i * 0.3), dxfattribs={"layer": "SCARI"})
        msp.add_line((x0, 0), (x0, 2.1), dxfattribs={"layer": "SCARI"})
    blk = doc.blocks.new("STAIR")
    for i in range(8):
        blk.add_line((0, i * 0.3), (1.2, i * 0.3))
    msp.add_blockref("STAIR", (40, 0), dxfattribs={"layer": "SCARI"})
    path = tmp_path / "stairs.dxf"
    doc.saveas(path)
    d = extract(path, "stairs.dxf", load_doc(path))
    plan = {"items": [item("st", "count_groups", [src("SCARI", ["line"])], "buc")]}
    plan["items"][0]["components"][0]["tolerance_m"] = 0.35
    res = execute_plan(plan, d, {})["items"][0]
    assert res["value"] == 3 and res["unit"] == "buc"
    # group counts add up with block counts in one item (both "buc")
    plan["items"][0]["components"].append(
        {
            "label": "bloc",
            "measure": "count",
            "sources": [src("SCARI", ["block"])],
            "subtract_sources": [],
            "region": None,
            "tolerance_m": None,
            "factors": [],
        }
    )
    assert execute_plan(plan, d, {})["items"][0]["value"] == 4
    ev = res["components"][0]["evidence"]
    assert len(ev) == 3 and all(e.get("geometry") for e in ev) and "1.20 × 2.10 m" in ev[0]["note"]


def test_units_from_car_blocks():
    from shapely.geometry import Point, box

    feats = [F("closed", 0, 0, 300, 200)]
    for i in range(3):  # three 4.5 x 1.8 cars drawn in metres in a file saved as millimetres
        h = f"C{i}"
        feats.append(
            SimpleNamespace(
                kind="block",
                nested=False,
                block="CAR2_TOP",
                handle=h,
                geom=Point(10 * i, 0),
                text=None,
                text_height=None,
            )
        )
        feats.append(
            SimpleNamespace(
                kind="closed",
                nested=True,
                block="CAR2_TOP",
                handle=h,
                geom=box(10 * i, 0, 10 * i + 4.5, 1.8),
                text=None,
                text_height=None,
            )
        )
    ui = infer_units(4, feats)
    assert ui.code == 6 and ui.source == "inferred" and "mașini" in " ".join(ui.evidence)


def test_calibrated_scale_override(garden):
    import copy

    from dwg_takeoff.cad import apply_units

    d = copy.deepcopy(garden)
    apply_units(d, {"scale": 0.5})  # e.g. a plan drawn at plot scale: 1 unit = 0.5 m
    assert d.unit_scale == 0.5 and d.unit_info["source"] == "override"
    plan = {"items": [item("lawn", "area", [src("0", ["hatch"], color="#3ca03c")])]}
    assert execute_plan(plan, d, {})["items"][0]["value"] == pytest.approx(200 * 0.25)


def test_excel_export_has_live_formulas(drawing):
    import io

    import openpyxl

    from dwg_takeoff.export import result_to_xlsx

    beton = item("beton", "area", [src("PAVAJ_8CM", ["hatch"])], "m³")
    beton["components"][0]["factors"] = [
        {"name": "grosime", "value": 0.15, "unit": "m", "origin": "assumed", "note": ""}
    ]
    plan = {"items": [item("pavaj8", "area", [src("PAVAJ_8CM", ["hatch"])]), beton]}
    res = execute_plan(plan, drawing, {})
    wb = openpyxl.load_workbook(io.BytesIO(result_to_xlsx(res, "x.dxf", {"units": "millimeters"})))
    rez, comp = wb["Rezumat"], wb["Componente"]
    assert rez["C3"].value.startswith("=SUMIF(Componente!")
    row = next(r for r in range(2, comp.max_row + 1) if comp.cell(r, 2).value == "beton")
    assert comp.cell(row, 5).value == pytest.approx(750) and comp.cell(row, 8).value == 0.15
    assert (
        comp.cell(row, 15).value
        == f'=E{row}*IF(H{row}="",1,H{row})*IF(J{row}="",1,J{row})*IF(L{row}="",1,L{row})*IF(N{row}="",1,N{row})'
    )
    assert "Info" in wb.sheetnames


def test_dynamic_block_instances_counted_under_their_name(tmp_path):
    """Each instance of a dynamic block is an anonymous block (*U..) pointing to its definition."""
    doc = ezdxf.new("R2018")
    door = doc.blocks.new("DOOR")
    door.add_line((0, 0), (0.9, 0))
    for i in range(3):
        anon = doc.blocks.new(f"*U{10 + i}")
        anon.add_line((0, 0), (0.8 + 0.1 * i, 0))
        anon.block_record.set_xdata("AcDbBlockRepBTag", [(1070, 1), (1005, door.block_record.dxf.handle)])
        doc.modelspace().add_blockref(f"*U{10 + i}", (5 * i, 0), dxfattribs={"layer": "USI"})
    path = tmp_path / "doors.dxf"
    doc.saveas(path)
    d = extract(path, "doors.dxf", load_doc(path))
    assert {f.block for f in d.features if f.kind == "block"} == {"DOOR"}
    plan = {"items": [item("usi", "count", [src("USI", ["block"], block="DOOR")], "buc")]}
    assert execute_plan(plan, d, {})["items"][0]["value"] == 3


def test_numbers_read_from_labels_are_not_measurements(drawing):
    it = item("t", "count", [src("TEXT", ["text"])], "m²")
    it["components"][0]["factors"] = [
        {"name": "suprafață declarată", "value": 13.44, "unit": "m²", "origin": "drawing", "note": ""}
    ]
    res = execute_plan({"items": [it]}, drawing, {})["items"][0]
    assert res["status"] == "assumption_based"
    assert any("Se numără texte" in w for w in res["components"][0]["warnings"])


def test_block_with_distant_base_point_is_placed_where_it_is_drawn(tmp_path):
    """Symbol libraries save blocks whose graphics sit far from the base point: the block counts
    inside the region where it is seen, not where its insertion point happens to be."""
    from shapely.geometry import box

    doc = ezdxf.new("R2018")
    doc.header["$INSUNITS"] = 6
    far = doc.blocks.new("TREE_FAR")
    far.add_circle((40, -40), 1)  # drawn 56 m away from the base point
    near = doc.blocks.new("DOOR")
    near.add_line((0, 0), (1, 0))
    near.add_arc((0, 0), 1, 0, 90)
    msp = doc.modelspace()
    for x in (2, 6):
        msp.add_blockref("TREE_FAR", (x - 40, 45), dxfattribs={"layer": "PLANT"})
    msp.add_blockref("DOOR", (3, 3), dxfattribs={"layer": "DOOR"})
    row = doc.blocks.new("ROW")  # nested: a row of two such trees
    for x in (0, 3):
        row.add_blockref("TREE_FAR", (x - 40, 40))
    msp.add_blockref("ROW", (1, 8), dxfattribs={"layer": "ROWS"})
    path = tmp_path / "far.dxf"
    doc.saveas(path)
    d = extract(path, "far.dxf", load_doc(path))
    blocks = {f.handle: f for f in d.features if f.kind == "block" and not f.nested}
    pts = sorted((round(f.geom.x, 3), round(f.geom.y, 3)) for f in blocks.values() if f.block == "TREE_FAR")
    assert pts == [(2.0, 5.0), (6.0, 5.0)]
    door = next(f for f in blocks.values() if f.block == "DOOR")
    assert (door.geom.x, door.geom.y) == (3, 3)  # inside its own graphics: insertion point kept
    plan = {"items": [item("trees", "count", [src("PLANT", ["block"])], unit="buc")]}
    plan["items"][0]["components"][0]["region"] = "site"
    res = execute_plan(plan, d, {"site": box(0, 0, 10, 10)})
    assert res["items"][0]["value"] == 2
    nested = sorted(
        (round(f.geom.x, 3), round(f.geom.y, 3))
        for f in d.features
        if f.kind == "block" and f.nested and f.block == "TREE_FAR"
    )
    assert nested == [(1.0, 8.0), (4.0, 8.0)]
    plan = {"items": [item("row_trees", "count", [src("ROWS", ["block"], block="TREE_FAR")], unit="buc")]}
    plan["items"][0]["components"][0]["region"] = "site"
    assert execute_plan(plan, d, {"site": box(0, 0, 10, 10)})["items"][0]["value"] == 2


def test_files_that_are_not_drawings_are_named(tmp_path):
    from dwg_takeoff.cad.convert import ConversionError, to_dxf

    page = tmp_path / "download.dwg"
    page.write_text("<!DOCTYPE html><html>403 Forbidden</html>")
    with pytest.raises(ConversionError, match="pagină web"):
        to_dxf(page)
    doc = ezdxf.new("R2018")
    doc.modelspace().add_line((0, 0), (1, 0))
    renamed = tmp_path / "really-a-dxf.dwg"
    doc.saveas(renamed)
    assert to_dxf(renamed) == renamed  # read as DXF, no converter needed
